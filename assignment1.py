#Ezekiel Tan, 241764L, Tut02
from click import style
from pandas.core.config_init import pc_pprint_nest_depth
import shelve
from tabulate import tabulate
from openpyxl import Workbook
from openpyxl.styles import Border, Side, Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import datetime
import re
from colorama import Fore, Style, init
import textwrap
init(autoreset=True)


class Employee:
    def __init__(self, name, id, email, programmes, department, status):
        self.__name = name
        self.__ID = id
        self.__email = email
        self.__programmes = programmes
        self.__department = department
        self.__status = status

    def set_name(self, name):
        self.__name = name
    def get_name(self):
        return self.__name
    def set_ID(self, id):
        self.__ID = id
    def get_ID(self):
        return self.__ID
    def set_email(self, email):
        self.__email = email
    def get_email(self):
        return self.__email
    def set_programmes(self, programmes):
        self.__programmes = programmes
    def get_programmes(self):
        return self.__programmes
    def set_department(self, department):
        self.__department = department
    def get_department(self):
        return self.__department
    def set_status(self, status):
        self.__status = status
    def get_status(self):
        return self.__status

    def add_programme(self, programme):
        new_normalized = normalize_programme_name(programme)
        for existing in self.__programmes:
            if normalize_programme_name(existing) == new_normalized:
                print(f"Programme '{Fore.YELLOW}{programme}{Style.RESET_ALL}' already exists for this employee.")
                return
        self.__programmes.append(programme.strip())
        print(f"Programme '{Fore.YELLOW}{programme.strip()}{Style.RESET_ALL}' added.")

        # Log the action
        with shelve.open("AdminLog", writeback=True) as log_db:
            logs = log_db.get("logs", [])
            logs.append({
                "timestamp": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "admin": admin_username,
                "action": f"Added programme '{Fore.YELLOW}{programme.strip()}{Style.RESET_ALL}' to employee ID {Fore.BLUE}{self.__ID}{Style.RESET_ALL}"
            })
            log_db["logs"] = logs

    def remove_programme(self, programme):
        target_normalized = normalize_programme_name(programme)
        for existing in self.__programmes:
            if normalize_programme_name(existing) == target_normalized:
                self.__programmes.remove(existing)
                print(f"Programme '{Fore.YELLOW}{existing}{Style.RESET_ALL}' removed.")
                with shelve.open("AdminLog", writeback=True) as log_db:
                    logs = log_db.get("logs", [])
                    logs.append({
                        "timestamp": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                        "admin": admin_username,
                        "action": f"Removed programme '{Fore.YELLOW}{existing}{Style.RESET_ALL}' from employee ID {Fore.BLUE}{self.__ID}{Style.RESET_ALL}"
                    })
                    log_db["logs"] = logs
                return
        print(f"Programme '{Fore.YELLOW}{programme.strip()}{Style.RESET_ALL}' not found in the employee's list.")

    def to_list(self):
        return [
            self.__ID,
            self.__name,
            self.__email,
            ', '.join(self.__programmes),
            self.__department,
            "Full-time" if self.__status else "Part-time"
        ]

    def __str__(self):
        program = ""
        employee_id = str(self.__ID)
        for prog in self.__programmes:
            program += prog + "\n"
        data = [[
            Fore.CYAN + employee_id + Style.RESET_ALL,
            Fore.GREEN + self.__name + Style.RESET_ALL,
            self.__email,
            program,
            self.__department,
            Fore.GREEN + "Full-time" + Style.RESET_ALL if self.__status else Fore.MAGENTA + "Part-time" + Style.RESET_ALL
        ]]
        headers = [
            Fore.YELLOW + "ID" + Style.RESET_ALL,
            Fore.YELLOW + "Name" + Style.RESET_ALL,
            Fore.YELLOW + "Email" + Style.RESET_ALL,
            Fore.YELLOW + "Enrolled Programmes" + Style.RESET_ALL,
            Fore.YELLOW + "Department" + Style.RESET_ALL,
            Fore.YELLOW + "Status" + Style.RESET_ALL
        ]
        table = tabulate(
            data,
            headers=headers,
            tablefmt="fancy_grid",
            colalign=("center", "left", "left", "left", "left", "center"),
            stralign="center"
        )

        # Add a title with dynamic border
        table_width = len(table.split('\n')[0])
        border = "═" * table_width
        title = "EMPLOYEE DETAILS".center(table_width)

        return f"\n{border}\n{Fore.YELLOW}{title}{Style.RESET_ALL}\n{border}\n{table}\n"

def admin_login():
    max_attempts = 3
    attempts = 0

    # username: admin1   password: password123
    # username: admin2   password: pass123
    # username: admin3   password: password000

    with shelve.open("Admin") as admin_db:
        accounts = admin_db.get("accounts", {})

    current_username = None

    while attempts < max_attempts:
        if current_username is None:
            attempts = 0
            username = input(f"Enter admin username (or '{Fore.BLUE}exit{Style.RESET_ALL}' to quit): ").strip()

            if username.lower() == 'exit':
                print("Exiting program.")
                exit()

            if username not in accounts:
                print(f"{Fore.YELLOW}Username not found.{Style.RESET_ALL}")
                continue

            current_username = username

        password = input(f"Password (or '{Fore.BLUE}back{Style.RESET_ALL}' to change user): ").strip()

        if password.lower() == "back":
            current_username = None
            continue
        # Verify password
        if accounts[current_username] == password:
            print(f"{Fore.GREEN}Login successful!{Style.RESET_ALL}\n")
            return current_username
        else:
            attempts += 1
            print(f"{Fore.YELLOW}Incorrect password.{Style.RESET_ALL} Attempts left: {max_attempts - attempts}")

    print("Too many failed attempts. Exiting program.")
    exit()

def normalize_programme_name(name):
    # Lowercase, strip outer spaces, and remove all internal non-alphanumeric characters
    return re.sub(r'\W+', '', name.strip().lower())

def is_valid_name(name):
    return bool(re.fullmatch(r"[A-Za-z ]{1,30}", name.strip()))

def is_valid_email(email):
    # Simple regex for email validation
    regex = r'^[\w\.-]+@[\w\.-]+\.\w+$'
    return bool(re.match(regex, email.strip()))

def manage_programme():
    with shelve.open("User", writeback=True) as db:
        users = db.get('employee', {})

        while True:
            employee_input = input(f"\nEnter Employee ID (type '{Fore.BLUE}cancel{Style.RESET_ALL}' to return): ")

            if employee_input.strip().lower() == "cancel":
                return

            try:
                employee_id = int(employee_input)
                break

            except ValueError:
                print(f"{Fore.YELLOW}Invalid ID.{Style.RESET_ALL} Please enter a numeric value.")

        if employee_id not in users:
            print(f"{Fore.RED}Employee not found.{Style.RESET_ALL}")
        else:
            employee = users[employee_id]
            print(employee)

            while True:
                action = input(
                    f"Do you want to ({Fore.YELLOW}a{Style.RESET_ALL})dd or ({Fore.YELLOW}r{Style.RESET_ALL})emove a programme? (type '{Fore.BLUE}cancel{Style.RESET_ALL}' to return): ").strip().lower()

                if action == 'cancel':
                    return  # exits to main menu
                elif action not in ['a', 'r']:
                    print(f"{Fore.YELLOW}Invalid option.{Style.RESET_ALL} Please enter '{Fore.BLUE}a{Style.RESET_ALL}', '{Fore.BLUE}r{Style.RESET_ALL}', or '{Fore.BLUE}cancel{Style.RESET_ALL}'.")
                    continue

                while True:
                    programme_name = input(f"Enter programme name (or type '{Fore.BLUE}cancel{Style.RESET_ALL}' to return): ").strip()

                    if programme_name.lower() == 'cancel':
                        return
                    elif programme_name == "":
                        print(f"{Fore.YELLOW}Programme name cannot be empty.{Style.RESET_ALL} Please enter a valid name or type '{Fore.BLUE}cancel{Style.RESET_ALL}'.")
                        continue
                    else:
                        break

                if action == 'a':
                    employee.add_programme(programme_name)
                elif action == 'r':
                    employee.remove_programme(programme_name)

                break

            # Save back
            users[employee_id] = employee
            db['employee'] = users
            print(employee)


def quicksort_employees(employee_list, low, high):
    """Recursive QuickSort implementation for employee names"""
    if low < high:
        # Partition the list
        pi = partition(employee_list, low, high)

        # Recursively sort elements before and after partition
        quicksort_employees(employee_list, low, pi - 1)
        quicksort_employees(employee_list, pi + 1, high)


def partition(employee_list, low, high):
    """Helper function for QuickSort"""
    pivot = employee_list[high].get_name().lower()
    i = low - 1

    for j in range(low, high):
        if employee_list[j].get_name().lower() <= pivot:
            i += 1
            employee_list[i], employee_list[j] = employee_list[j], employee_list[i]

    employee_list[i + 1], employee_list[high] = employee_list[high], employee_list[i + 1]
    return i + 1


def merge_sort_employees(employee_list):
    """Merge sort implementation for employees with dual ascending criteria"""
    if len(employee_list) > 1:
        mid = len(employee_list) // 2
        left = employee_list[:mid]
        right = employee_list[mid:]

        merge_sort_employees(left)
        merge_sort_employees(right)

        i = j = k = 0

        while i < len(left) and j < len(right):
            # Primary sort: programme count (ascending)
            left_prog = len(left[i][1].get_programmes())
            right_prog = len(right[j][1].get_programmes())

            if left_prog < right_prog:
                employee_list[k] = left[i]
                i += 1
            elif left_prog > right_prog:
                employee_list[k] = right[j]
                j += 1
            else:
                # Secondary sort: employee ID (ascending)
                if left[i][0] < right[j][0]:
                    employee_list[k] = left[i]
                    i += 1
                else:
                    employee_list[k] = right[j]
                    j += 1
            k += 1

        while i < len(left):
            employee_list[k] = left[i]
            i += 1
            k += 1

        while j < len(right):
            employee_list[k] = right[j]
            j += 1
            k += 1

def export_user_data_to_excel(admin_username):
    with shelve.open("User") as db:
        users = db.get('employee', {})

        if not users:
            print(f"{Fore.RED}No employee data found to export.{Style.RESET_ALL}")
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "Employee Records"

        headers = ["ID", "Name", "Email", "Enrolled Programme", "Department", "Status"]
        ws.append(headers)

        # Style definitions
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        center_align = Alignment(vertical='center', horizontal='center', wrap_text=True)
        header_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
        header_font = Font(bold=True)

        fill_colors = [
            PatternFill(start_color="D9EAF7", end_color="D9EAF7", fill_type="solid"),  # Light blue
            PatternFill(start_color="EDEDED", end_color="EDEDED", fill_type="solid")   # Light gray
        ]

        # Style the header
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border

        row_index = 2
        color_toggle = 0

        for emp in users.values():
            programmes = emp.get_programmes() or [""]
            num_rows = len(programmes)

            # Append rows per enrolled programme
            for i, programme in enumerate(programmes):
                ws.append([
                    emp.get_ID() if i == 0 else "",
                    emp.get_name() if i == 0 else "",
                    emp.get_email() if i == 0 else "",
                    programme,
                    emp.get_department() if i == 0 else "",
                    "Full-time" if emp.get_status() else "Part-time" if i == 0 else ""
                ])

            # Merge identical user cells
            if num_rows > 1:
                for col in [1, 2, 3, 5, 6]:
                    ws.merge_cells(start_row=row_index, start_column=col,
                                   end_row=row_index + num_rows - 1, end_column=col)

            # Apply consistent fill and borders for this user block
            fill = fill_colors[color_toggle % 2]
            for r in range(row_index, row_index + num_rows):
                for c in range(1, 7):
                    cell = ws.cell(row=r, column=c)
                    cell.border = thin_border
                    cell.alignment = center_align
                    cell.fill = fill

            # Toggle fill color for next user
            color_toggle += 1
            row_index += num_rows

        # Auto-fit column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value)) if cell.value else 0 for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max_len + 2

        # Add filter to header row
        ws.auto_filter.ref = f"A1:F{row_index - 1}"

        filename = "employee_data_export.xlsx"
        wb.save(filename)
        print(f"\nExcel file '{Fore.YELLOW}{filename}{Style.RESET_ALL}' generated successfully!")

    # Log the action
    with shelve.open("AdminLog", writeback=True) as log_db:
        logs = log_db.get("logs", [])
        logs.append({
            "timestamp": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "admin": admin_username,
            "action": "Exported user data to Excel"
        })
        log_db["logs"] = logs

def main_menu(admin_username):
    while True:
        print(f"""
{Fore.BLUE}╔════════════════════════════╗
{Fore.BLUE}{Fore.YELLOW}          MAIN MENU          {Fore.BLUE}
{Fore.BLUE}╠════════════════════════════╣

{Fore.BLUE}{Fore.CYAN} a.{Style.RESET_ALL} Display all employee records {Fore.BLUE}
{Fore.BLUE}{Fore.CYAN} b.{Style.RESET_ALL} Add a new employee record   {Fore.BLUE}
{Fore.BLUE}{Fore.CYAN} c.{Style.RESET_ALL} Enroll in training program  {Fore.BLUE}
{Fore.BLUE}{Fore.CYAN} d.{Style.RESET_ALL} Sort by Department         {Fore.BLUE}
{Fore.BLUE}{Fore.CYAN} e.{Style.RESET_ALL} Sort by Program Count      {Fore.BLUE}
{Fore.BLUE}{Fore.CYAN} f.{Style.RESET_ALL} Search employee            {Fore.BLUE}
{Fore.BLUE}{Fore.CYAN} g.{Style.RESET_ALL} Exit program               {Fore.BLUE}
{Fore.BLUE}{Fore.CYAN} h.{Style.RESET_ALL} Export to Excel            {Fore.BLUE}
{Fore.BLUE}{Fore.CYAN} i.{Style.RESET_ALL} View admin logs            {Fore.BLUE}

{Fore.BLUE}╚════════════════════════════╝
{Style.RESET_ALL}""")

        choice = input("Please choose an option (A-I): ").strip().lower()
        if choice == "a":
            with shelve.open("User") as db:
                users = db.get('employee', {})
                if not users:
                    print(Fore.RED + "\nNo employee records found." + Style.RESET_ALL)
                else:
                    headers = [
                        Fore.YELLOW + "S/No" + Style.RESET_ALL,
                        Fore.YELLOW + "ID" + Style.RESET_ALL,
                        Fore.YELLOW + "Name" + Style.RESET_ALL,
                        Fore.YELLOW + "Email" + Style.RESET_ALL,
                        Fore.YELLOW + "Programme(s)" + Style.RESET_ALL,
                        Fore.YELLOW + "Department" + Style.RESET_ALL,
                        Fore.YELLOW + "Status" + Style.RESET_ALL
                    ]

                    table_data = []

                    for idx, emp in enumerate(users.values(), 1):
                        # Format programmes with each on new line
                        programmes = emp.get_programmes()
                        if programmes:
                            program = '\n'.join([f"• {p}" for p in programmes])
                        else:
                            program = "N/A"

                        table_data.append([
                            Fore.CYAN + str(idx) + Style.RESET_ALL,
                            emp.get_ID(),
                            Fore.GREEN + emp.get_name() + Style.RESET_ALL,
                            emp.get_email(),
                            program,  # This will show as multiple lines
                            emp.get_department(),
                            Fore.GREEN + "Full-time" + Style.RESET_ALL if emp.get_status() else Fore.MAGENTA + "Part-time" + Style.RESET_ALL
                        ])

                    table_str = tabulate(
                        table_data,
                        headers=headers,
                        tablefmt="fancy_grid",
                        colalign=("center", "center", "left", "left", "left", "left", "center"),
                        stralign="center"
                    )

                    # Calculate table width
                    table_width = len(table_str.split('\n')[0])

                    # Print the header with dynamic border
                    print("\n" + "═" * table_width)
                    print(Fore.YELLOW + "EMPLOYEE RECORDS".center(table_width) + Style.RESET_ALL)
                    print("═" * table_width)

                    # Print the table
                    print(table_str)

            input("\nPress Enter to return to menu:")

        elif choice == "b":
            numOfEmployee = input(f"How many employees would you like to add? (or type '{Fore.BLUE}cancel{Style.RESET_ALL}' to return): ").strip()
            if numOfEmployee.lower() == "cancel":
                continue
            if not numOfEmployee.isdigit() or int(numOfEmployee) <= 0:
                print(f"{Fore.YELLOW}Please enter a positive number.{Style.RESET_ALL}")
                input("\nPress Enter to return to menu:")
                continue
            numOfEmployee = int(numOfEmployee)

            for n in range(numOfEmployee):
                with shelve.open("User") as db:
                    users = db.get('employee', {})

                # ID input and validation
                while True:
                    id_input = input(f"Enter ID (max 5 digits) (or type '{Fore.BLUE}cancel{Style.RESET_ALL}' to return): ").strip()
                    if id_input.lower() == "cancel":
                        break
                    if not id_input.isdigit():
                        print(f"{Fore.YELLOW}Invalid input.{Style.RESET_ALL} Please enter numeric digits only.")
                        continue
                    if len(id_input) > 5:
                        print(f"{Fore.YELLOW}ID cannot exceed 5 digits.{Style.RESET_ALL}")
                        continue
                    id = int(id_input)
                    if id in users:
                        print(f"ID {Fore.BLUE}{id}{Style.RESET_ALL} already exists. Please enter a different ID.")
                    else:
                        break
                else:
                    continue

                if id_input.lower() == "cancel":
                    break

                # Name input and validation
                while True:
                    name = input(
                        f"Enter name (alphabets and spaces only, max 30 characters) (or type '{Fore.BLUE}cancel{Style.RESET_ALL}' to return): ").strip()
                    if name.lower() == "cancel":
                        break
                    if not name:
                        print(f"{Fore.YELLOW}Name cannot be empty.{Style.RESET_ALL}")
                        continue
                    if len(name) > 30:
                        print(f"{Fore.YELLOW}Name cannot exceed 30 characters.{Style.RESET_ALL}")
                        continue
                    if not all(char.isalpha() or char.isspace() for char in name):
                        print(f"{Fore.YELLOW}Name can only contain alphabets and spaces.{Style.RESET_ALL}")
                        continue
                    break
                else:
                    continue

                if name.lower() == "cancel":
                    break

                # Email input and validation
                while True:
                    email = input(f"Enter email (valid format) (or '{Fore.BLUE}cancel{Style.RESET_ALL}'): ").strip()
                    if email.lower() == "cancel":
                        break
                    if not email:
                        print(f"{Fore.YELLOW}Email cannot be empty.{Style.RESET_ALL}")
                        continue
                    if not is_valid_email(email):
                        print(f"{Fore.YELLOW}Invalid email format.{Style.RESET_ALL}")
                        continue
                    # Check email uniqueness
                    email_exists = any(emp.get_email().lower() == email.lower() for emp in users.values())
                    if email_exists:
                        print(f"{Fore.YELLOW}Email already exists.{Style.RESET_ALL} Enter a different email.")
                        continue
                    break
                else:
                    continue

                if email.lower() == "cancel":
                    break

                # Department input (no empty)
                while True:
                    department = input(f"Enter department (or '{Fore.BLUE}cancel{Style.RESET_ALL}'): ").strip()
                    if department.lower() == "cancel":
                        break
                    if not department:
                        print(f"{Fore.YELLOW}Department cannot be empty.{Style.RESET_ALL}")
                        continue
                    break
                else:
                    continue

                if department.lower() == "cancel":
                    break

                # Status input (True/False)
                while True:
                    status_input = input(
                        f"Enter working status ({Fore.YELLOW}True{Style.RESET_ALL} for Full-time, {Fore.YELLOW}False{Style.RESET_ALL} for Part-time) (or '{Fore.BLUE}cancel{Style.RESET_ALL}'): ").strip()
                    if status_input.lower() == "cancel":
                        break
                    if status_input.lower() == "true":
                        status = True
                        break
                    elif status_input.lower() == "false":
                        status = False
                        break
                    else:
                        print(f"{Fore.YELLOW}Invalid input.{Style.RESET_ALL} Please enter {Fore.GREEN}True{Style.RESET_ALL} or {Fore.GREEN}False{Style.RESET_ALL} else {Fore.BLUE}cancel{Style.RESET_ALL}.")
                else:
                    continue

                if status_input.lower() == "cancel":
                    break

                # Programmes list is empty on add
                programmes = []

                # Log action
                with shelve.open("AdminLog", writeback=True) as log_db:
                    logs = log_db.get("logs", [])
                    logs.append({
                        "timestamp": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                        "admin": admin_username,
                        "action": f"Added employee: {Fore.GREEN}{name}{Style.RESET_ALL} (ID: {Fore.BLUE}{id}{Style.RESET_ALL})"
                    })
                    log_db["logs"] = logs

                # Save employee
                with shelve.open("User") as db:
                    users = db.get('employee', {})
                    empData = Employee(name, id, email, programmes, department, status)
                    users[id] = empData
                    db['employee'] = users

                print("\nNew employee added:")
                print(empData)

            input("Press enter to return to the menu:")

        elif choice == "c":
            manage_programme()

            input("Press enter to return to the menu:")

        elif choice == "d":
            with shelve.open("User") as db:
                users = db.get('employee', {})

                if not users:
                    print(Fore.RED + "\nNo employee records found." + Style.RESET_ALL)
                else:
                    departments = {}
                    for emp in users.values():
                        dept = emp.get_department().strip().lower()
                        departments.setdefault(dept, []).append(emp)

                    # Sort departments alphabetically
                    sorted_departments = sorted(departments.keys())

                    for dept in sorted_departments:
                        employees = departments[dept]

                        # QuickSort employees by name (case-insensitive)
                        if employees:  # Only sort if list is non-empty
                            quicksort_employees(employees, 0, len(employees) - 1)

                        # Prepare table data
                        table_data = []
                        for idx, emp in enumerate(employees, 1):
                            programmes = emp.get_programmes()
                            formatted_programmes = '\n'.join([f"• {p}" for p in programmes]) if programmes else "N/A"
                            table_data.append([
                                Fore.YELLOW + str(idx) + Style.RESET_ALL,
                                emp.get_ID(),
                                Fore.GREEN + emp.get_name() + Style.RESET_ALL,
                                emp.get_email(),
                                formatted_programmes,
                                emp.get_department(),
                                Fore.GREEN + "Full-time" + Style.RESET_ALL if emp.get_status() else Fore.MAGENTA + "Part-time" + Style.RESET_ALL
                            ])

                        # Print department header
                        dept_header = f" {dept.upper()} DEPARTMENT ({len(employees)} employees) "
                        print(f"\n{'═' * len(dept_header)}")
                        print(Fore.RED + dept_header + Style.RESET_ALL)
                        print(f"{'═' * len(dept_header)}")

                        # Print table
                        headers = [
                            Fore.CYAN + "S/No" + Style.RESET_ALL,
                            Fore.CYAN + "ID" + Style.RESET_ALL,
                            Fore.CYAN + "Name" + Style.RESET_ALL,
                            Fore.CYAN + "Email" + Style.RESET_ALL,
                            Fore.CYAN + "Programmes" + Style.RESET_ALL,
                            Fore.CYAN + "Department" + Style.RESET_ALL,
                            Fore.CYAN + "Status" + Style.RESET_ALL
                        ]
                        print(tabulate(table_data, headers=headers, tablefmt="fancy_grid"))

            input("Press enter to return to the menu:")

        elif choice == "e":
            with shelve.open("User") as db:
                users = db.get('employee', {})
                if not users:
                    print(Fore.RED + "\nNo employee records found." + Style.RESET_ALL)
                else:
                    user_list = list(users.items())

                    # Perform merge sort
                    merge_sort_employees(user_list)

                    # Prepare table data
                    table_data = []
                    for idx, (emp_id, emp) in enumerate(user_list, start=1):
                        programmes = emp.get_programmes()
                        formatted_programmes = '\n'.join(
                            [f"• {p}" for p in programmes]) if programmes else Fore.YELLOW + "N/A" + Style.RESET_ALL

                        table_data.append([
                            Fore.CYAN + str(idx) + Style.RESET_ALL,
                            emp_id,
                            Fore.GREEN + emp.get_name() + Style.RESET_ALL,
                            emp.get_email(),
                            formatted_programmes,
                            emp.get_department(),
                            Fore.GREEN + "Full-time" + Style.RESET_ALL if emp.get_status() else Fore.MAGENTA + "Part-time" + Style.RESET_ALL,
                            len(programmes)
                        ])

                    headers = [
                        Fore.YELLOW + "S/No" + Style.RESET_ALL,
                        Fore.YELLOW + "ID" + Style.RESET_ALL,
                        Fore.YELLOW + "Name" + Style.RESET_ALL,
                        Fore.YELLOW + "Email" + Style.RESET_ALL,
                        Fore.YELLOW + "Programmes" + Style.RESET_ALL,
                        Fore.YELLOW + "Department" + Style.RESET_ALL,
                        Fore.YELLOW + "Status" + Style.RESET_ALL,
                        Fore.YELLOW + "Total" + Style.RESET_ALL
                    ]

                    table_str = tabulate(
                        table_data,
                        headers=headers,
                        tablefmt="fancy_grid",
                        colalign=("center", "center", "left", "left", "left", "left", "center", "center")
                    )

                    # Dynamic header
                    table_width = len(table_str.split('\n')[0])
                    title = "EMPLOYEES SORTED BY PROGRAMME COUNT & ID (ASCENDING)"
                    print(f"\n{Fore.BLUE}{'═' * table_width}{Style.RESET_ALL}")
                    print(Fore.BLUE + title.center(table_width) + Style.RESET_ALL)
                    print(f"{Fore.BLUE}{'═' * table_width}{Style.RESET_ALL}")
                    print(table_str)

            input("Press enter to return to the menu:")

        elif choice == "f":
            query = input(f"\nEnter {Fore.GREEN}ID{Style.RESET_ALL} or {Fore.GREEN}name{Style.RESET_ALL} prefix to search (type '{Fore.BLUE}cancel{Style.RESET_ALL}' to return to menu): ").strip()
            if query.lower() == "cancel":
                continue
            print(f"\nSearching for: {Fore.GREEN}{query}{Style.RESET_ALL}")

            with shelve.open("User") as db:
                users = db.get('employee', {})
                results = []

                for emp_id, emp in users.items():
                    # Convert ID to string for matching
                    id_str = str(emp_id)
                    name = emp.get_name()

                    # Check for matches (ID prefix or name prefix)
                    if (id_str.startswith(query) or
                            name.lower().startswith(query.lower())):
                        # Format programmes with bullet points
                        programmes = emp.get_programmes()
                        formatted_programmes = '\n'.join([f"• {p}" for p in programmes]) if programmes else "N/A"

                        results.append([
                            Fore.CYAN + id_str + Style.RESET_ALL,
                            Fore.GREEN + name + Style.RESET_ALL,
                            emp.get_email(),
                            formatted_programmes,
                            emp.get_department(),
                            Fore.GREEN + "Full-time" + Style.RESET_ALL if emp.get_status() else Fore.MAGENTA + "Part-time" + Style.RESET_ALL
                        ])

            # Display results
            if results:
                headers = [
                    Fore.YELLOW + "ID" + Style.RESET_ALL,
                    Fore.YELLOW + "Name" + Style.RESET_ALL,
                    Fore.YELLOW + "Email" + Style.RESET_ALL,
                    Fore.YELLOW + "Programmes" + Style.RESET_ALL,
                    Fore.YELLOW + "Department" + Style.RESET_ALL,
                    Fore.YELLOW + "Status" + Style.RESET_ALL
                ]

                print(f"\n{Fore.BLUE}=== Found {len(results)} matching employees ==={Style.RESET_ALL}")
                print(tabulate(
                    results,
                    headers=headers,
                    tablefmt="fancy_grid",
                    colalign=("center", "left", "left", "left", "left", "center")
                ))
            else:
                print(f"\n{Fore.RED}No matching employees found.{Style.RESET_ALL}")

            input("Press enter to return to the menu:")

        elif choice == "g":
            break

        elif choice == "h":
            export_user_data_to_excel(admin_username)
            input("Press enter to return to the menu:")

        elif choice == "i":
            with shelve.open("AdminLog") as log_db:
                logs = log_db.get("logs", [])

                if not logs:
                    print(Fore.YELLOW + "\nNo logs available." + Style.RESET_ALL)
                else:
                    # Prepare table data with ALL entries
                    table_data = []
                    for log in logs:  # Removed [-20:] to show all entries
                        timestamp = log.get("timestamp", "")
                        try:
                            dt = datetime.fromisoformat(timestamp)
                            # Split date and time with different colors
                            date_part = Fore.CYAN + dt.strftime("%Y-%m-%d") + Style.RESET_ALL
                            time_part = Fore.MAGENTA + dt.strftime("%H:%M:%S") + Style.RESET_ALL
                            formatted_time = f"{date_part} {time_part}"
                        except:
                            formatted_time = Fore.RED + timestamp + Style.RESET_ALL

                        table_data.append([
                            formatted_time,
                            Fore.GREEN + log.get("admin", "N/A") + Style.RESET_ALL,
                            log.get("action", "No action recorded")
                        ])

                    # Generate table first to determine width
                    headers = [
                        Fore.YELLOW + "Timestamp" + Style.RESET_ALL,
                        Fore.YELLOW + "Admin" + Style.RESET_ALL,
                        Fore.YELLOW + "Action" + Style.RESET_ALL
                    ]

                    table_str = tabulate(
                        table_data,
                        headers=headers,
                        tablefmt="fancy_grid",
                        colalign=("center", "left", "left")
                    )

                    # Calculate table width
                    table_width = len(table_str.split('\n')[0])

                    # Create dynamic header (removed "showing last 20")
                    header_text = f" ADMIN LOGS ({len(logs)} entries) "
                    border = "═" * table_width
                    print(f"\n{Fore.BLUE}{border}{Style.RESET_ALL}")
                    print(Fore.BLUE + header_text.center(table_width) + Style.RESET_ALL)
                    print(f"{Fore.BLUE}{border}{Style.RESET_ALL}")

                    # Print the table
                    print(table_str)

            input("Press enter to return to the menu:")


if __name__ == "__main__":
    admin_username = admin_login()
    if admin_username:
        main_menu(admin_username)