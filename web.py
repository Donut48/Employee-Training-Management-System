#Ezekiel Tan 241764L Tut_02
from models import RequestWrapper
from flask import Flask, jsonify, render_template, send_file, request, session, redirect, url_for, current_app
from flask_socketio import SocketIO, join_room, emit, disconnect
from cryptography.fernet import Fernet
import json
import shelve
import re
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Border, Side, Alignment, Font
from openpyxl.utils import get_column_letter
from fpdf import FPDF
import io
import uuid
import heapq
from datetime import datetime
import bcrypt
from functools import total_ordering, wraps
from typing import List
import random
import string
import smtplib
from email.mime.text import MIMEText
import bisect
import os
import traceback
import hashlib
from werkzeug.security import generate_password_hash, check_password_hash
from dotenv import load_dotenv

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY")
socketio = SocketIO(app)


cipher = Fernet(os.getenv("FERNET_KEY"))

email_address = os.getenv("EMAIL_ADDRESS")
email_password = os.getenv("EMAIL_PASSWORD")

class Employee:
    def __init__(self, name, id, email, programmes, department, status):
        self.__name = name
        self.__ID = int(id)
        self.__email = email
        self.__programmes = programmes
        self.__department = department
        self.__status = status

    def set_name(self, name):
        self.__name = name
    def get_name(self):
        return self.__name
    def set_ID(self, id):
        self.__ID = id
    def get_ID(self):
        return self.__ID
    def set_email(self, email):
        self.__email = email
    def get_email(self):
        return self.__email
    def set_programmes(self, programmes):
        self.__programmes = programmes
    def get_programmes(self):
        return self.__programmes
    def set_department(self, department):
        self.__department = department
    def get_department(self):
        return self.__department
    def set_status(self, status):
        self.__status = status
    def get_status(self):
        return self.__status

    def add_programme(self, programme, admin_username):
        """Add a programme to the employee's list"""
        if not programme or not programme.strip():
            return {'success': False, 'error': 'Programme name cannot be empty'}

        programme = programme.strip()
        new_normalized = normalize_name(programme)
        for existing in self.__programmes:
            if normalize_name(existing) == new_normalized:
                return {
                    'success': False,
                    'error': 'Programme already exists',
                    'existing': existing,
                    'new': programme
                }

        self.__programmes.append(programme)

        # Log the action
        with shelve.open("AdminLog", writeback=True) as log_db:
            logs = log_db.get("logs", [])
            logs.append({
                "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "admin": admin_username,
                "action": f"Added programme '{programme}' to employee ID {self.__ID}",
                "employee_id": self.__ID,
                "programme": programme
            })
            log_db["logs"] = logs

        return {'success': True, 'message': f'Programme "{programme}" added'}

    def remove_programme(self, programme, admin_username):
        """Remove a programme from the employee's list"""
        if not programme or not programme.strip():
            return {'success': False, 'error': 'Programme name cannot be empty'}

        target_normalized = normalize_name(programme)
        for existing in self.__programmes:
            if normalize_name(existing) == target_normalized:
                self.__programmes.remove(existing)

                # Log the action
                with shelve.open("AdminLog", writeback=True) as log_db:
                    logs = log_db.get("logs", [])
                    logs.append({
                        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                        "admin": admin_username,
                        "action": f"Removed programme '{existing}' from employee ID {self.__ID}",
                        "employee_id": self.__ID,
                        "programme": existing
                    })
                    log_db["logs"] = logs

                return {'success': True, 'message': f'Programme "{existing}" removed'}

        return {'success': False, 'error': f'Programme "{programme.strip()}" not found'}

    def to_list(self):
        return [
            self.__ID,
            self.__name,
            self.__email,
            ', '.join(self.__programmes),
            self.__department,
            "Full-time" if self.__status else "Part-time"
        ]

    @staticmethod
    def from_dict(data):
        return Employee(
            name=data['name'],
            id=data['id'],
            email=data['email'],
            programmes=data.get('programmes', []),
            department=data['department'],
            status=data['status'] == 'Full-time'  # Convert to boolean
        )

    def to_dict(self):
        return {
            'name': self.__name,
            'id': self.__ID,
            'email': self.__email,
            'programmes': self.__programmes,
            'department': self.__department,
            'status': 'Full-time' if self.__status else 'Part-time'
        }

def admin_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'username' not in session:
            return jsonify({'error': 'Unauthorized'}), 401
        return f(*args, **kwargs)
    return decorated_function

def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'employee_id' not in session and 'username' not in session:
            return jsonify({'error': 'Unauthorized'}), 401
        return f(*args, **kwargs)
    return decorated

def encrypt_data(data):
    return cipher.encrypt(json.dumps(data).encode()).decode()

def decrypt_data(encrypted_data):
    return json.loads(cipher.decrypt(encrypted_data.encode()).decode())

def get_chat_key(participants):
    return hashlib.md5(json.dumps(sorted(participants)).encode()).hexdigest()

def hash_password(password):
    return generate_password_hash(password)

def check_password(hashed_password, password):
    return check_password_hash(hashed_password, password)

def update_chat_last_read(chat_key, user_id):
    with shelve.open('ChatDB', writeback=True) as db:
        if chat_key in db:
            chat_data = decrypt_data(db[chat_key])

            # Validate structure length
            if len(chat_data) < 4:
                # Not expected structure
                return

            # The list of participant dicts starts at index 1 up to len(chat_data) - 4
            # (Because last three indices are last_read, chat_type, calls)
            participant_dicts = chat_data[1: -3]

            # Find the dict index that contains user_id
            user_msgs = None
            for participant_dict in participant_dicts:
                if user_id in participant_dict:
                    user_msgs = participant_dict[user_id]
                    break

            if user_msgs is None:
                # user_id not found in chat participants
                return

            # Now update last read position
            last_read_list = chat_data[-3]  # list of dicts [{user_id: last_read_index}, ...]

            # Find the entry for this user_id
            for read_status in last_read_list:
                if user_id in read_status:
                    # Update last read index to last message index (or -1 if no messages)
                    last_index = len(user_msgs) - 1 if len(user_msgs) > 0 else -1
                    read_status[user_id] = last_index
                    break
            else:
                # If user not found in last_read_list, add a new entry
                last_index = len(user_msgs) - 1 if len(user_msgs) > 0 else -1
                last_read_list.append({user_id: last_index})

            # Encrypt and save back
            db[chat_key] = encrypt_data(chat_data)


@app.route('/')
@app.route('/index')
def index():
    is_employee = session.get('is_employee', False)
    is_admin = session.get('is_admin', False)

    print("CURRENT SESSION:", dict(session))

    # Always use home.html, but pass auth status
    return render_template('home.html',
                           is_admin=is_admin,
                           is_employee=is_employee,
                           employee_id=session.get('employee_id'))


@app.route('/api/employee-data')
def get_employee_data():
    try:
        with shelve.open('User') as db:
            data = {
                'total_employees': 0,
                'employment_types': {'Full-time': 0, 'Part-time': 0},
                'programmes': {},
                'programme_display_names': {},  # Stores original programme names
                'departments': {},
                'department_display_names': {},  # Stores original department names
                'employees': {}
            }

            # Helper function to track original names
            def track_original_name(normalized, original, display_names):
                if normalized not in display_names:
                    display_names[normalized] = original
                elif len(original) < len(display_names[normalized]):
                    # Prefer shorter names for display
                    display_names[normalized] = original

            # Get all employee records
            employees = {}
            for key in db.keys():
                item = db[key]
                if isinstance(item, Employee):
                    employees[key] = item
                elif isinstance(item, dict):
                    if key.lower() in ['employees', 'employee']:
                        employees.update(item)
                    else:
                        employees[key] = item

            # Process each employee
            for emp_id, emp_obj in employees.items():
                try:
                    # Convert Employee object to dictionary
                    if isinstance(emp_obj, Employee):
                        emp_data = {
                            'id': emp_obj.get_ID(),
                            'name': emp_obj.get_name(),
                            'email': emp_obj.get_email(),
                            'programmes': emp_obj.get_programmes(),
                            'department': emp_obj.get_department(),
                            'status': emp_obj.get_status()
                        }
                    else:
                        emp_data = emp_obj

                    # Store employee data
                    data['employees'][emp_id] = emp_data

                    # Count employment types
                    status = 'Full-time' if emp_data.get('status', False) else 'Part-time'
                    data['employment_types'][status] += 1

                    # Process department (case-insensitive counting)
                    dept = emp_data.get('department', 'Unknown')
                    if dept:
                        norm_dept = normalize_name(dept)
                        data['departments'][norm_dept] = data['departments'].get(norm_dept, 0) + 1
                        track_original_name(norm_dept, dept, data['department_display_names'])

                    # Process programmes (case-insensitive counting)
                    programmes = emp_data.get('programmes', [])
                    for programme in (programmes if isinstance(programmes, list) else [programmes]):
                        if programme:
                            norm_prog = normalize_name(programme)
                            data['programmes'][norm_prog] = data['programmes'].get(norm_prog, 0) + 1
                            track_original_name(norm_prog, programme, data['programme_display_names'])

                    data['total_employees'] += 1

                except Exception as e:
                    print(f"Error processing employee {emp_id}: {str(e)}")
                    continue

            return jsonify(data)

    except Exception as e:
        print(f"Database error: {str(e)}")
        return jsonify({
            'error': True,
            'message': str(e),
            'total_employees': 0,
            'employment_types': {'Full-time': 0, 'Part-time': 0},
            'programmes': {},
            'programme_display_names': {},
            'departments': {},
            'department_display_names': {},
            'employees': {}
        })

def normalize_name(name):
    if not name:
        return ''
    return re.sub(r'[^a-z0-9]', '', str(name).strip().lower())

@app.route('/allRecords')
def allRecords():

    return render_template('allRecords.html')


@app.route('/api/export-excel-direct')
def export_excel_direct():
    try:
        username = session.get('username', 'system')

        with shelve.open('User') as db:
            if 'employee' not in db:
                return "No employee data found to export", 404

            employees = db['employee']
            if not employees:
                return "Employee data is empty", 404

            output = io.BytesIO()
            wb = Workbook()
            ws = wb.active
            ws.title = "Employee Records"

            # Set column widths
            ws.column_dimensions['A'].width = 15  # ID
            ws.column_dimensions['B'].width = 25  # Name
            ws.column_dimensions['C'].width = 30  # Email
            ws.column_dimensions['D'].width = 30  # Programmes
            ws.column_dimensions['E'].width = 25  # Department
            ws.column_dimensions['F'].width = 15  # Status

            headers = ["Employee ID", "Full Name", "Email Address", "Enrolled Programmes", "Department",
                       "Employment Status"]
            ws.append(headers)

            # Create styles
            header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)
            even_row_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
            odd_row_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
            border = Border(left=Side(style='thin'), right=Side(style='thin'),
                            top=Side(style='thin'), bottom=Side(style='thin'))
            center_align = Alignment(vertical='center', horizontal='center', wrap_text=True)
            left_align = Alignment(vertical='center', horizontal='left')

            # Style headers
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.border = border
                cell.alignment = center_align

            row_index = 2
            color_toggle = 0  # Track alternating colors across all rows

            for emp_id, emp_obj in employees.items():
                # Get employee data
                if hasattr(emp_obj, 'get_ID'):
                    emp_data = {
                        'id': emp_obj.get_ID(),
                        'name': emp_obj.get_name(),
                        'email': emp_obj.get_email(),
                        'programmes': emp_obj.get_programmes(),
                        'department': emp_obj.get_department(),
                        'status': emp_obj.get_status()
                    }
                else:
                    emp_data = emp_obj

                programmes = emp_data.get('programmes', []) if isinstance(emp_data, dict) else getattr(emp_obj,
                                                                                                       'programmes', [])
                if not isinstance(programmes, list):
                    programmes = [programmes] if programmes else []

                # Determine fill color for this employee block
                fill = even_row_fill if color_toggle % 2 == 0 else odd_row_fill
                num_rows = max(1, len(programmes))  # At least one row per employee

                if not programmes:
                    # Single row for employees with no programmes
                    ws.append([
                        emp_data.get('id', ''),
                        emp_data.get('name', ''),
                        emp_data.get('email', ''),
                        'Not enrolled',
                        emp_data.get('department', ''),
                        "Full-time" if (
                            emp_data.get('status') if isinstance(emp_data, dict) else getattr(emp_obj, 'status',
                                                                                              False)) else "Part-time"
                    ])

                    # Style the row
                    for cell in ws[row_index]:
                        cell.border = border
                        cell.fill = fill
                        cell.alignment = left_align if cell.column_letter in ['B', 'C', 'D', 'E'] else center_align

                    row_index += 1
                else:
                    # Multiple rows for employees with programmes
                    for i, programme in enumerate(programmes):
                        ws.append([
                            emp_data.get('id', '') if i == 0 else "",
                            emp_data.get('name', '') if i == 0 else "",
                            emp_data.get('email', '') if i == 0 else "",
                            programme,
                            emp_data.get('department', '') if i == 0 else "",
                            "Full-time" if (
                                emp_data.get('status') if isinstance(emp_data, dict) else getattr(emp_obj, 'status',
                                                                                                  False)) else "Part-time" if i == 0 else ""
                        ])

                        # Style the row
                        for cell in ws[row_index]:
                            cell.border = border
                            cell.fill = fill
                            cell.alignment = left_align if cell.column_letter in ['B', 'C', 'D', 'E'] else center_align

                        row_index += 1

                    # Merge cells for employee info
                    if len(programmes) > 1:
                        for col in ['A', 'B', 'C', 'E', 'F']:
                            ws.merge_cells(start_row=row_index - len(programmes),
                                           start_column=ord(col) - 64,
                                           end_row=row_index - 1,
                                           end_column=ord(col) - 64)

                # Toggle color for next employee
                color_toggle += 1

            # Add auto-filter
            ws.auto_filter.ref = ws.dimensions

            # Freeze header row
            ws.freeze_panes = 'A2'

            wb.save(output)
            output.seek(0)

            with shelve.open("AdminLog", writeback=True) as log_db:
                logs = log_db.get("logs", [])
                logs.append({
                    "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    "admin": username,
                    "action": "Exported employee data to Excel",
                    "file": "employee_records.xlsx",
                    "record_count": len(employees)
                })
                log_db["logs"] = logs

            return send_file(
                output,
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                as_attachment=True,
                download_name='employee_records.xlsx'
            )

    except Exception as e:
        print(f"Excel export error: {str(e)}")
        return jsonify({'error': True, 'message': str(e)}), 500


@app.route('/api/export-pdf-direct')
def export_pdf_direct():
    try:
        username = session.get('username', 'system')

        with shelve.open('User') as db:
            if 'employee' not in db:
                return "No employee data found to export", 404

            employees = db['employee']
            if not employees:
                return "Employee data is empty", 404

            pdf = FPDF(orientation='L')  # Landscape orientation for better fit
            pdf.add_page()
            pdf.set_font("Arial", size=10)

            # Add title
            pdf.set_font("Arial", 'B', 16)
            pdf.cell(0, 10, "Employee Records", 0, 1, 'C')
            pdf.ln(5)
            pdf.set_font("Arial", size=10)

            # Column widths (adjusted for landscape)
            col_widths = [20, 30, 50, 40, 30, 15, 70]  # Added serial number column

            # Colors
            header_color = (79, 129, 189)  # Blue
            row_color = (220, 230, 241)  # Light blue
            white = (255, 255, 255)

            # Headers
            headers = ["No.", "Employee ID", "Full Name", "Email", "Department", "Status", "Programmes"]
            pdf.set_fill_color(*header_color)
            pdf.set_text_color(255, 255, 255)
            pdf.set_font('Arial', 'B', 10)

            for i, header in enumerate(headers):
                pdf.cell(col_widths[i], 10, header, 1, 0, 'C', True)
            pdf.ln()

            # Data rows
            pdf.set_font('Arial', '', 10)
            pdf.set_text_color(0, 0, 0)

            for count, (emp_id, emp_obj) in enumerate(employees.items(), 1):
                # Get employee data (same as before)
                if hasattr(emp_obj, 'get_ID'):
                    emp_data = {
                        'id': emp_obj.get_ID(),
                        'name': emp_obj.get_name(),
                        'email': emp_obj.get_email(),
                        'programmes': emp_obj.get_programmes(),
                        'department': emp_obj.get_department(),
                        'status': emp_obj.get_status()
                    }
                else:
                    emp_data = emp_obj

                programmes = emp_data.get('programmes', []) if isinstance(emp_data, dict) else getattr(emp_obj,
                                                                                                       'programmes', [])
                if not isinstance(programmes, list):
                    programmes = [programmes] if programmes else []

                # Alternate row colors
                fill = row_color if count % 2 == 0 else white
                pdf.set_fill_color(*fill)

                # If no programmes, display "Not enrolled"
                if not programmes:
                    pdf.cell(col_widths[0], 10, str(count), 'LR', 0, 'C', True)
                    pdf.cell(col_widths[1], 10, str(emp_data.get('id', '')), 'LR', 0, 'C', True)
                    pdf.cell(col_widths[2], 10, str(emp_data.get('name', '')), 'LR', 0, 'L', True)
                    pdf.cell(col_widths[3], 10, str(emp_data.get('email', '')), 'LR', 0, 'L', True)
                    pdf.cell(col_widths[4], 10, str(emp_data.get('department', '')), 'LR', 0, 'L', True)
                    pdf.cell(col_widths[5], 10, "Full-time" if (
                        emp_data.get('status') if isinstance(emp_data, dict) else getattr(emp_obj, 'status',
                                                                                          False)) else "Part-time",
                             'LR', 0, 'C', True)
                    pdf.cell(col_widths[6], 10, "Not enrolled", 'LR', 1, 'L', True)
                else:
                    # First programme in first row
                    pdf.cell(col_widths[0], 10, str(count), 'LR', 0, 'C', True)
                    pdf.cell(col_widths[1], 10, str(emp_data.get('id', '')), 'LR', 0, 'C', True)
                    pdf.cell(col_widths[2], 10, str(emp_data.get('name', '')), 'LR', 0, 'L', True)
                    pdf.cell(col_widths[3], 10, str(emp_data.get('email', '')), 'LR', 0, 'L', True)
                    pdf.cell(col_widths[4], 10, str(emp_data.get('department', '')), 'LR', 0, 'L', True)
                    pdf.cell(col_widths[5], 10, "Full-time" if (
                        emp_data.get('status') if isinstance(emp_data, dict) else getattr(emp_obj, 'status',
                                                                                          False)) else "Part-time",
                             'LR', 0, 'C', True)
                    pdf.cell(col_widths[6], 10, str(programmes[0]), 'LR', 1, 'L', True)

                    # Additional programmes in subsequent rows
                    for programme in programmes[1:]:
                        pdf.set_fill_color(*fill)
                        pdf.cell(col_widths[0], 10, "", 'LR', 0, 'C', True)
                        pdf.cell(col_widths[1], 10, "", 'LR', 0, 'C', True)
                        pdf.cell(col_widths[2], 10, "", 'LR', 0, 'L', True)
                        pdf.cell(col_widths[3], 10, "", 'LR', 0, 'L', True)
                        pdf.cell(col_widths[4], 10, "", 'LR', 0, 'L', True)
                        pdf.cell(col_widths[5], 10, "", 'LR', 0, 'C', True)
                        pdf.cell(col_widths[6], 10, str(programme), 'LR', 1, 'L', True)

                # Add bottom border for the last row of each employee
                pdf.set_fill_color(*fill)
                for i in range(len(col_widths)):
                    pdf.cell(col_widths[i], 0, "", 'T', 0, 'C', True)
                pdf.ln()

            # Add document footer
            pdf.set_y(-15)
            pdf.set_font('Arial', 'I', 8)
            pdf.cell(0, 10, f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", 0, 0, 'C')

            # Save to memory
            output = io.BytesIO()
            pdf_bytes = pdf.output(dest='S').encode('latin-1')
            output.write(pdf_bytes)
            output.seek(0)

            with shelve.open("AdminLog", writeback=True) as log_db:
                logs = log_db.get("logs", [])
                logs.append({
                    "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                    "admin": username,
                    "action": "Exported employee data to PDF",
                    "file": "employee_records.pdf",
                    "record_count": len(employees)
                })
                log_db["logs"] = logs

            return send_file(
                output,
                mimetype='application/pdf',
                as_attachment=True,
                download_name='employee_records.pdf'
            )

    except Exception as e:
        print(f"PDF export error: {str(e)}")
        return jsonify({'error': True, 'message': str(e)}), 500

@app.route('/manageEmp')
def manageEmp():
    return render_template('manageEmp.html')

def is_valid_name(name):
    return bool(re.fullmatch(r"[A-Za-z ]{1,30}", name.strip()))

def is_valid_email(email):
    regex = r'^[\w\.-]+@[\w\.-]+\.\w+$'
    return bool(re.match(regex, email.strip()))

def is_valid_id(id):
    return isinstance(id, int) and 0 < id <= 99999

def _id_exists(id):
    with shelve.open('User') as db:
        for key in db.keys():
            item = db[key]
            if isinstance(item, Employee) and item.get_ID() == id:
                return True
            elif isinstance(item, dict) and item.get('id') == id:
                return True
    return False

def _email_exists(email):
    email = email.lower()
    with shelve.open('User') as db:
        for key in db.keys():
            item = db[key]
            if isinstance(item, Employee) and item.get_email().lower() == email:
                return True
            elif isinstance(item, dict) and str(item.get('email', '')).lower() == email:
                return True
    return False


def generate_password():
    """Generate a random 12-character password meeting standard requirements"""
    upper = random.choices(string.ascii_uppercase, k=2)
    lower = random.choices(string.ascii_lowercase, k=3)
    digits = random.choices(string.digits, k=3)
    special = random.choices('!@#$%^&*', k=2)
    remaining = random.choices(string.ascii_letters + string.digits, k=2)
    password = upper + lower + digits + special + remaining
    random.shuffle(password)
    return ''.join(password)


def send_credentials(email, username, password, is_reset=False):
    """Send login credentials to employee's email"""
    try:
        # Email content
        if is_reset:
            subject = 'Your Password Has Been Reset'
            body = f"""
            Your password has been reset:
            Username: {username}
            New Password: {password}

            Please change your password after logging in.
            """
        else:
            subject = 'Your Employee Account Credentials'
            body = f"""
            Your employee account has been created:
            Username: {username}
            Temporary Password: {password}

            Please change your password after first login.
            """

        msg = MIMEText(body)
        msg['Subject'] = subject
        msg['From'] = 'no-reply@yourcompany.com'  # Change this to your sending email
        msg['To'] = email

        # SMTP Configuration
        with smtplib.SMTP('smtp.gmail.com', 587) as server:
            server.starttls()
            server.login(email_address, email_password)
            server.send_message(msg)

        print(f"✅ Email sent to {email}")
        return True

    except Exception as e:
        print(f"❌ Failed to send email: {e}")
        return False


@app.route('/api/employees', methods=['POST'])
@app.route('/api/employees/<int:employee_id>', methods=['PUT'])
@admin_required
def save_employee(employee_id=None):
    try:
        data = request.get_json()
        if not data:
            return jsonify({'success': False, 'error': 'No data received'}), 400

        admin_username = session.get('username', 'system')

        # Validate and convert ID to integer
        try:
            emp_id = int(data['id'])
            if len(str(emp_id)) != 5:
                return jsonify({'success': False, 'error': 'ID must be 5 digits'}), 400
        except (ValueError, KeyError):
            return jsonify({'success': False, 'error': 'Invalid employee ID'}), 400

        # Validate required fields
        required_fields = ['id', 'name', 'email', 'status']
        if not all(field in data for field in required_fields):
            return jsonify({'success': False, 'error': 'Missing required fields'}), 400

        # Initialize variables
        username = str(emp_id)
        email_sent = False
        password = generate_password()
        hashed_pw = bcrypt.hashpw(password.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')

        try:
            # Database operations in a single context
            with shelve.open('User', writeback=True) as user_db, \
                    shelve.open('Admin', writeback=True) as admin_db, \
                    shelve.open("AdminLog", writeback=True) as log_db:

                # Initialize employee storage if not exists
                user_db.setdefault('employee', {})
                admin_db.setdefault('employees', {})
                log_db.setdefault('logs', [])

                if request.method == 'PUT':
                    if employee_id != emp_id:
                        return jsonify({'success': False, 'error': 'ID in URL and body must match'}), 400

                    if emp_id not in user_db['employee']:
                        return jsonify({'success': False, 'error': 'Employee not found'}), 404

                    # Update existing employee
                    employee = user_db['employee'][emp_id]

                    # Check if email is being changed to an existing one
                    if data['email'] != employee.get_email():  # Changed to use getter
                        if any(e.get_email() == data['email'] for e in
                               user_db['employee'].values()):  # Changed to use getter
                            return jsonify({'success': False, 'error': 'Email already in use'}), 400

                    # Update fields
                    employee.set_name(data['name'])
                    employee.set_email(data['email'])
                    employee.set_department(data.get('department', ''))
                    employee.set_status(data['status'] == 'Full-time')

                    # Add log entry for update
                    log_db['logs'].append({
                        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                        "admin": admin_username,
                        "action": f"Updated employee: {data['name']} (ID: {emp_id})",
                        "employee_id": emp_id,
                        "email": data['email'],
                        "department": data.get('department', '')
                    })

                    # For creates (POST request)
                else:
                    if emp_id in user_db['employee']:
                        return jsonify({
                            'success': False,
                            'error': f'Employee ID {emp_id} already exists'
                        }), 400

                    if any(e.get_email() == data['email'] for e in
                           user_db['employee'].values()):  # Changed to use getter
                        return jsonify({'success': False, 'error': 'Email already in use'}), 400

                    # Create new employee
                    user_db['employee'][emp_id] = Employee(
                        name=data['name'],
                        id=emp_id,
                        email=data['email'],
                        programmes=[],
                        department=data.get('department', ''),
                        status=data['status'] == 'Full-time'
                    )

                    # Store account credentials (only for new employees)
                    admin_db['employees'][emp_id] = {
                        'id': emp_id,
                        'username': str(emp_id),
                        'password': hashed_pw
                    }

                    # Add log entry for creation
                    log_db['logs'].append({
                        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                        "admin": admin_username,
                        "action": f"Added employee: {data['name']} (ID: {emp_id})",
                        "employee_id": emp_id,
                        "email": data['email'],
                        "department": data.get('department', '')
                    })

                # Sync all databases
                user_db.sync()
                admin_db.sync()
                log_db.sync()

        except Exception as db_error:
            print(f"Database error: {db_error}")
            return jsonify({
                'success': False,
                'error': 'Database operation failed',
                'debug': str(db_error)
            }), 500

        # Email sending outside database context
        if request.method == "POST":
            try:
                send_credentials(data['email'], str(emp_id), password)
                email_sent = True
            except Exception as email_error:
                print(f"Email failed: {str(email_error)}")

        return jsonify({
            'success': True,
            'message': f'Employee { "updated" if request.method == "PUT" else "created" } successfully',
            'data': {
                'id': emp_id,
                'name': data['name'],
                'email': data['email'],
                'department': data.get('department', ''),
                'status': data['status']
            }
        }), 201 if request.method == 'POST' else 200

    except Exception as e:
        print(f"Unexpected error: {str(e)}")
        return jsonify({
            'success': False,
            'error': 'An unexpected error occurred',
            'debug': str(e)
        }), 500

def get_all_employees():
    employees = []
    with shelve.open('User') as db:
        for key in db.keys():
            item = db[key]
            if isinstance(item, Employee):
                employees.append(item)
            elif isinstance(item, dict) and 'id' in item:
                # Convert dict to Employee if needed
                employees.append(Employee(**item))
    return employees

@app.route('/api/check-employee-id')
def check_employee_id():
    try:
        id = request.args.get('id', '').strip()
        print(f"DEBUG: Checking ID: {id} (type: {type(id)})")

        if not id:
            return jsonify({'error': 'ID parameter required'}), 400

        if not id.isdigit() or len(id) > 5:
            return jsonify({'isUnique': False, 'reason': 'Invalid ID format'}), 400

        with shelve.open('User') as db:
            if 'employee' not in db:
                return jsonify({'isUnique': True})

            employees = db['employee']
            # Convert all keys to strings for consistent comparison
            employee_ids = [str(k) for k in employees.keys()]
            print(f"DEBUG: All employee IDs (as strings): {employee_ids}")

            if id in employee_ids:
                print(f"DEBUG: ID {id} exists (exact string match)")
                return jsonify({'isUnique': False})

            # Additional check for integer matches if needed
            try:
                if int(id) in employees:
                    print(f"DEBUG: ID {id} exists (integer match)")
                    return jsonify({'isUnique': False})
            except ValueError:
                pass

            print(f"DEBUG: ID {id} is truly available")
            return jsonify({'isUnique': True})

    except Exception as e:
        print(f"ERROR checking ID: {str(e)}")
        return jsonify({'error': str(e)}), 500


@app.route('/api/check-employee-email')
def check_employee_email():
    try:
        email = request.args.get('email', '').strip().lower()
        print(f"DEBUG: Checking email: {email}")

        if not email:
            return jsonify({'error': 'Email parameter required'}), 400

        if not re.match(r'^[\w\.-]+@[\w\.-]+\.\w+$', email):
            return jsonify({'isUnique': False, 'reason': 'Invalid email format'}), 400

        with shelve.open('User') as db:
            if 'employee' not in db:
                return jsonify({'isUnique': True})

            employees = db['employee']
            print(f"DEBUG: Checking against {len(employees)} employees")

            for emp_id, employee in employees.items():
                # Handle all possible formats:
                # 1. Employee object with get_email() method
                if hasattr(employee, 'get_email'):
                    if employee.get_email().lower() == email:
                        print(f"DEBUG: Found match in Employee object (ID: {emp_id})")
                        return jsonify({'isUnique': False})

                # 2. List format [name, email, ...]
                elif isinstance(employee, list) and len(employee) > 1:
                    if str(employee[1]).lower() == email:
                        print(f"DEBUG: Found match in list format (ID: {emp_id})")
                        return jsonify({'isUnique': False})

                # 3. Dictionary format {'email': ...}
                elif isinstance(employee, dict):
                    if str(employee.get('email', '')).lower() == email:
                        print(f"DEBUG: Found match in dict format (ID: {emp_id})")
                        return jsonify({'isUnique': False})

            print("DEBUG: Email is available")
            return jsonify({'isUnique': True})

    except Exception as e:
        print(f"ERROR checking email: {str(e)}")
        return jsonify({'error': str(e)}), 500

@app.route('/api/employees/<int:employee_id>', methods=['DELETE'])
@admin_required
def delete_employee(employee_id):
    try:
        admin_username = session.get('username', 'system')
        employee_name = None

        # Open all databases in a single context
        with shelve.open('User', writeback=True) as user_db, \
             shelve.open('Admin', writeback=True) as admin_db, \
             shelve.open("AdminLog", writeback=True) as log_db:

            # Initialize databases if needed
            user_db.setdefault('employee', {})
            admin_db.setdefault('employees', {})
            log_db.setdefault('logs', [])

            # Check if employee exists
            if employee_id not in user_db['employee']:
                return jsonify({
                    'success': False,
                    'error': 'Employee not found'
                }), 404

            # Get employee name for logging
            employee_obj = user_db['employee'][employee_id]
            employee_name = employee_obj.get_name()

            # Delete employee records
            del user_db['employee'][employee_id]
            if employee_id in admin_db['employees']:
                del admin_db['employees'][employee_id]

            # Add deletion log
            log_db['logs'].append({
                "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                "admin": admin_username,
                "action": f"Deleted employee: {employee_name} (ID: {employee_id})",
                "employee_id": employee_id
            })

            # Force sync all changes
            user_db.sync()
            admin_db.sync()
            log_db.sync()

        return jsonify({
            'success': True,
            'message': f'Employee {employee_name} (ID: {employee_id}) deleted successfully',
            'deleted_id': employee_id
        }), 200

    except Exception as e:
        print(f"Error deleting employee: {str(e)}")
        traceback.print_exc()
        return jsonify({
            'success': False,
            'error': 'Failed to delete employee',
            'debug': str(e)
        }), 500

@app.route('/api/employees/<int:employee_id>/programmes', methods=['POST'])
@admin_required
def add_employee_programme(employee_id):
    try:
        data = request.get_json()
        if not data or 'programme' not in data:
            return jsonify({'success': False, 'error': 'Programme name required'}), 400

        with shelve.open('User', writeback=True) as db:
            if 'employee' not in db or employee_id not in db['employee']:
                return jsonify({'success': False, 'error': 'Employee not found'}), 404

            employee = db['employee'][employee_id]
            result = employee.add_programme(data['programme'], session.get('username', 'system'))
            db.sync()

            if result['success']:
                return jsonify(result), 200
            return jsonify(result), 400

    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/employees/<int:employee_id>/programmes', methods=['DELETE'])
@admin_required
def remove_employee_programme(employee_id):
    try:
        data = request.get_json()
        if not data or 'programme' not in data:
            return jsonify({'success': False, 'error': 'Programme name required'}), 400

        with shelve.open('User', writeback=True) as db:
            if 'employee' not in db or employee_id not in db['employee']:
                return jsonify({'success': False, 'error': 'Employee not found'}), 404

            employee = db['employee'][employee_id]
            result = employee.remove_programme(data['programme'], session.get('username', 'system'))
            db.sync()

            if result['success']:
                return jsonify(result), 200
            return jsonify(result), 400

    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/logs')
def logs():
    return render_template('logs.html')


@app.route('/api/admin-logs')
def get_admin_logs():
    try:
        with shelve.open('AdminLog') as db:
            logs = db.get('logs', [])

            if session.get('is_employee'):
                employee_id = int(session.get('employee_id'))  # store as int for comparison
                filtered_logs = []

                for log in logs:
                    # Case 1: employee performed the action (as admin)
                    if str(log.get('admin')) == str(employee_id):
                        filtered_logs.append(log)
                        continue

                    # Case 2: employee is the target of the action
                    if log.get('employee_id') == employee_id:
                        filtered_logs.append(log)
                        continue

                    # Case 3: action text mentions them (like "(ID: 10005)")
                    if f"(ID: {employee_id})" in log.get('action', ''):
                        filtered_logs.append(log)

                logs = filtered_logs

            return jsonify({
                'success': True,
                'logs': logs[-1000:],  # cap at 1000
                'count': len(logs)
            })

    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/request-logs')
def get_request_logs():
    try:
        with shelve.open('AdminLog') as db:
            logs = db.get('request_logs', [])

            # Employees: filter to their own requests
            if session.get('is_employee'):
                employee_id = str(session.get('employee_id'))
                logs = [log for log in logs if str(log.get('employee_id')) == employee_id]

            formatted_logs = []
            for log in logs:
                formatted_logs.append({
                    'timestamp': log.get('timestamp'),
                    'admin': log.get('admin', 'system'),
                    'action': log.get('action', 'unknown').capitalize(),
                    'employee_id': log.get('employee_id', 'N/A'),
                    'request_type': log.get('request_type', 'Unknown'),
                    'details': log.get('details', 'No details available'),
                    'request_id': log.get('request_id', '')
                })

            return jsonify({
                'success': True,
                'logs': formatted_logs[-1000:],  # Return most recent 1000
                'count': len(formatted_logs)
            })

    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/requests')
def requests():
    return render_template('request.html')

class EmployeeRequest:
    def __init__(self, employee_id, request_type, priority, details):
        self.id = str(uuid.uuid4())
        self.employee_id = employee_id
        self.request_type = request_type
        self.priority = int(priority)
        self.details = details
        self.timestamp = datetime.datetime.now()

    def __lt__(self, other):
        if self.priority == other.priority:
            return self.timestamp < other.timestamp
        return self.priority < other.priority


def init_databases():
    """Initialize databases with proper queue synchronization"""
    with shelve.open('requests_db') as db:
        if 'pending_requests' not in db:
            db['pending_requests'] = {}
            db['request_queue'] = []

        # Ensure queue matches pending requests
        pending = db.get('pending_requests', {})
        current_queue = db.get('request_queue', [])

        # Rebuild queue if out of sync
        if len(pending) != len(current_queue):
            new_queue = []
            for req_id, request in pending.items():
                heapq.heappush(new_queue, RequestWrapper(request['priority'], request))
            db['request_queue'] = new_queue


init_databases()


@app.route('/employee-requests')
def employee_requests():
    # Get session data
    employee_id = session.get('employee_id', None)
    is_employee = session.get('is_employee', False)

    # Send these to the template
    return render_template(
        'employee_requests.html',
        employee_id=employee_id,
        is_employee=is_employee
    )

def employee_exists(employee_id, db):
    """Helper function for clearer logic"""
    sorted_ids = sorted(db['employees'].keys())
    idx = bisect.bisect_left(sorted_ids, employee_id)
    return idx < len(sorted_ids) and sorted_ids[idx] == employee_id

@app.route('/api/validate-employee/<employee_id>')
def validate_employee(employee_id):
    try:
        with shelve.open('Admin') as db:
            return jsonify({'exists': employee_exists(employee_id, db)})
    except Exception as e:
        return jsonify({'error': str(e)}), 500

def init_requests_db():
    with shelve.open('requests_db') as db:
        if 'request_queue' not in db:
            db['request_queue'] = []


@app.route('/api/requests', methods=['POST'])
def add_request():
    try:
        data = request.json
        request_id = str(uuid.uuid4())

        # Auto-fill employee ID if needed
        if session.get('is_employee'):
            data['employee_id'] = session['employee_id']

        # Create request with new UUID
        new_request = {
            'id': request_id,
            'employee_id': data['employee_id'],
            'request_type': data['request_type'],
            'priority': int(data['priority']),
            'details': data['details'],
            'timestamp': datetime.now().isoformat(),
            'status': 'pending',
            'status_history': [{
                'status': 'pending',
                'timestamp': datetime.now().isoformat()
            }]
        }

        # Store using the wrapper class
        with shelve.open('requests_db', writeback=True) as db:
            # Update dictionary storage
            pending = db.get('pending_requests', {})
            pending[request_id] = new_request
            db['pending_requests'] = pending

            # Update priority queue with wrapper
            queue = db.get('request_queue', [])
            heapq.heappush(queue, RequestWrapper(new_request['priority'], new_request))
            db['request_queue'] = queue

            print(f"Added request {request_id} to database") #debug
        return jsonify({'success': True, 'request': new_request})

    except Exception as e:
        print(f"Error adding request: {str(e)}") #debug
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/requests/process', methods=['POST'])
def process_request():
    try:
        if not session.get('is_admin'):
            return jsonify({'success': False, 'error': 'Admin required'}), 403

        data = request.json
        if 'request_id' not in data or 'action' not in data:
            return jsonify({'success': False, 'error': 'Missing fields'}), 400

        with shelve.open('requests_db') as db:
            requests = db.get('requests', {})
            request_data = requests.pop(data['request_id'], None)
            if not request_data:
                return jsonify({'success': False, 'error': 'Request not found'}), 404
            db['requests'] = requests

        # Add to processed logs
        with shelve.open('AdminLog') as log_db:
            processed = log_db.get('processed_requests', {})
            request_data['status'] = data['action']  # 'approved' or 'denied'
            request_data['processed_at'] = datetime.now().isoformat()
            processed[request_data['id']] = request_data
            log_db['processed_requests'] = processed

        return jsonify({'success': True})

    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/requests/process-next', methods=['GET'])
def process_next_request():
    try:
        with shelve.open('requests_db') as db:
            request_queue = db.get('request_queue', [])

            if not request_queue:
                return jsonify({'success': False, 'error': 'Queue is empty'}), 404

            # Sort the queue by priority (ascending) and timestamp (ascending)
            request_queue.sort(key=lambda x: (x.request['priority'], x.request['timestamp']))

            # Get the highest priority request (lowest priority number)
            next_request = request_queue.pop(0)  # Pop the highest priority (sorted first)

            # Remove from pending_requests as well
            pending_requests = db.get('pending_requests', {})
            if next_request.request['id'] in pending_requests:
                del pending_requests[next_request.request['id']]

            db['pending_requests'] = pending_requests
            db['request_queue'] = request_queue

            # Log the processing
            log_request(next_request.request, "PROCESSED")

            return jsonify({
                'success': True,
                'request': next_request.request,
                'remaining': len(request_queue)
            })

    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/requests', methods=['GET'])
def get_requests():
    try:
        with shelve.open('requests_db') as db:
            pending = list(db.get('pending_requests', {}).values())
            processed = list(db.get('processed_requests', {}).values())

            # Apply filters
            filter_type = request.args.get('type')
            filter_status = request.args.get('status')
            filter_priority = request.args.get('priority')

            requests = pending + processed if filter_status == 'all' else pending

            filtered = [
                r for r in requests
                if (not filter_type or r['request_type'] == filter_type) and
                   (not filter_status or r['status'] == filter_status) and
                   (not filter_priority or str(r['priority']) == filter_priority)
            ]

            # Sort by priority then timestamp
            filtered.sort(key=lambda x: (x['priority'], x['timestamp']))

            return jsonify({
                'success': True,
                'requests': filtered,
                'pending_count': len(pending),
                'processed_count': len(processed)
            })

    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

def check_employee_exists(employee_id):
    """Consolidated employee validation"""
    with shelve.open('Admin') as db:
        employee_key1 = f"emp{employee_id}"
        employee_key2 = employee_id  # Handle both formats
        employees = db.get('employees', {})
        return employee_key1 in employees or employee_key2 in employees


def log_request(request_data, action, request_id=None, new_status=None):
    with shelve.open("AdminLog", writeback=True) as log_db:
        logs = log_db.get("request_logs", [])
        log_entry = {
            "timestamp": datetime.now().isoformat(),
            "request_id": request_id or request_data.get('id'),
            "employee_id": request_data['employee_id'],  # Use employee_id directly
            "request_type": request_data['request_type'],
            "priority": request_data['priority'],
            "action": action,
            "admin": session.get('username', 'system'),
            "status": new_status or request_data.get('status', 'pending'),
            "details": request_data.get('details', ''),
            "ip_address": request.remote_addr if hasattr(request, 'remote_addr') else None
        }

        log_entry = {k: v for k, v in log_entry.items() if v is not None}

        logs.append(log_entry)
        log_db["request_logs"] = logs[-1000:]

@app.route('/api/requests/stats')
def request_stats():
    try:
        with shelve.open('requests_db') as db:
            requests = list(db.get('requests', {}).values())

        if not requests:
            return jsonify({
                'success': True,
                'message': 'No pending requests'
            })

        # Basic counts
        stats = {
            'total_requests': len(requests),
            'count_by_type': {},
            'count_by_priority': {i: 0 for i in range(1, 6)},
            'oldest_request': min(r['timestamp'] for r in requests),
            'newest_request': max(r['timestamp'] for r in requests)
        }

        # Detailed breakdowns
        for req in requests:
            # Count by type
            stats['count_by_type'][req['request_type']] = \
                stats['count_by_type'].get(req['request_type'], 0) + 1

            # Count by priority
            stats['count_by_priority'][req['priority']] += 1

        # Time calculations
        now = datetime.now()
        stats['avg_wait_time'] = str(
            sum((now - datetime.fromisoformat(r['timestamp'])).total_seconds()
                for r in requests) / len(requests)
        )

        return jsonify({
            'success': True,
            'stats': stats
        })

    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@app.route('/api/requests/<request_id>/process', methods=['POST'])
def process_specific_request(request_id):
    try:
        if not session.get('is_admin'):
            return jsonify({'success': False, 'error': 'Admin access required'}), 403

        action = request.json.get('action')
        valid_actions = ['approve', 'deny', 'return']
        if action not in valid_actions:
            return jsonify({'success': False, 'error': 'Invalid action'}), 400

        with shelve.open('requests_db') as db:
            # 1. Remove from pending requests
            pending = db.get('pending_requests', {})
            request_data = pending.pop(request_id, None)

            if not request_data:
                return jsonify({'success': False, 'error': 'Request not found'}), 404

            # 2. Remove from priority queue
            queue = db.get('request_queue', [])
            queue = [rw for rw in queue if rw.request['id'] != request_id]

            # 3. Update request status based on action
            new_status = {
                'approve': 'approved',
                'deny': 'denied',
                'return': 'returned'
            }[action]

            request_data.update({
                'status': new_status,
                'processed_by': session.get('username', 'admin'),
                'processed_at': datetime.now().isoformat(),
                'status_history': request_data.get('status_history', []) + [{
                    'status': new_status,
                    'admin': session.get('username', 'admin'),
                    'timestamp': datetime.now().isoformat()
                }]
            })

            # 4. Store in appropriate location
            if action == 'return':
                # Return to employee - keep in pending but with updated status
                pending[request_id] = request_data
                heapq.heappush(queue, RequestWrapper(request_data['priority'], request_data))
            else:
                # Move to processed requests
                processed = db.get('processed_requests', {})
                processed[request_id] = request_data
                db['processed_requests'] = processed

            # 5. Save all changes
            db['pending_requests'] = pending
            db['request_queue'] = queue

            # 6. Add minimal log entry
            log_request(request_data, action, request_id=request_id, new_status=new_status)

        return jsonify({'success': True})

    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@app.route('/api/chats', methods=['GET'])
@login_required
def get_chats():
    # Identify user type and ID
    if session.get('is_employee'):
        user_id = session.get('employee_id')  # integer
    elif session.get('is_admin'):
        user_id = session.get('username')     # string
    else:
        return jsonify({'success': False, 'message': 'User not logged in'}), 401

    with shelve.open('ChatDB') as db:
        user_chats = []

        for key, encrypted_data in db.items():
            try:
                chat_data = decrypt_data(encrypted_data)

                if user_id in chat_data['participants']:
                    # Calculate unread messages
                    last_read = chat_data['read_positions'].get(str(user_id), -1)
                    total_messages = sum(len(m) for m in chat_data['messages'].values())
                    unread = max(0, total_messages - last_read - 1)

                    if chat_data['type'] == 'individual':
                        other_id = next(p for p in chat_data['participants'] if p != user_id)
                        with shelve.open('User') as user_db:
                            other_user = user_db.get(str(other_id), {})
                        chat_info = {
                            'key': key,
                            'last_activity': chat_data['last_activity'],
                            'unread': unread,
                            'type': 'individual',
                            'name': other_user.get('name', f'User {other_id}'),
                            'picture': other_user.get('profile_pic', '/static/images/default_profile.png')
                        }
                    else:
                        chat_info = {
                            'key': key,
                            'last_activity': chat_data['last_activity'],
                            'unread': unread,
                            'type': 'group',
                            'name': chat_data.get('group_name', f"Group {key[:6]}"),
                            'picture': chat_data.get('group_pic', '/static/images/default_group.png')
                        }
                    user_chats.append(chat_info)

            except Exception as e:
                print(f"Error processing chat {key}: {str(e)}")

        # Sort by most recent
        user_chats.sort(key=lambda x: x['last_activity'], reverse=True)
        return jsonify({'chats': user_chats})

def get_current_user_id():
    return session.get('user_id')

@app.route('/api/messages/<chat_key>')
def fetch_chat_messages(chat_key):
    user_id = session.get('user_id') or session.get('username')
    if not user_id:
        return jsonify({'error': 'Unauthorized'}), 401

    with shelve.open('Chat') as db:
        if chat_key not in db:
            return jsonify({'messages': []})

        encrypted_data = db[chat_key]
        chat_data = decrypt_data(encrypted_data)  # Assuming you use encryption

        messages = []
        if isinstance(chat_data, list) and len(chat_data) >= 2:
            all_user_messages = chat_data[1]  # Dict of user_id: list of messages
            for sender, message_list in all_user_messages.items():
                for item in message_list:
                    for time, content in item.items():
                        messages.append({
                            'sender_id': sender,
                            'timestamp': time,
                            'content': content
                        })

        # Sort messages by timestamp
        messages.sort(key=lambda m: m['timestamp'])
        return jsonify({'messages': messages})


@app.route('/api/chats/<chat_key>/messages', methods=['GET'])
@login_required
def get_chat_messages(chat_key):
    current_user_id = str(session.get('employee_id') or session.get('username'))

    # Use ChatDB consistently
    with shelve.open('ChatDB') as db:
        if chat_key not in db:
            return jsonify({'success': False, 'messages': []})

        chat_data = decrypt_data(db[chat_key])
        messages = []

        # Unified message extraction
        if 'messages' in chat_data:
            for sender_id, message_list in chat_data['messages'].items():
                username = get_display_name(sender_id)
                profile_pic = get_profile_pic(sender_id)  # Implement similar to get_display_name

                for msg in message_list:
                    if isinstance(msg, (list, tuple)) and len(msg) >= 2:
                        messages.append({
                            'sender': sender_id,
                            'text': msg[1],
                            'time': msg[0],
                            'is_me': str(sender_id) == current_user_id,
                            'username': username,
                            'profile_pic': profile_pic
                        })

        return jsonify({
            'success': True,
            'messages': sorted(messages, key=lambda x: x['time']),
            'chat_type': chat_data.get('type', 'individual')
        })


def get_profile_pic(user_id):
    """Get profile picture URL for a user"""
    try:
        with shelve.open('User') as udb:
            employee = udb.get('employees', {}).get(str(user_id))
            if employee and 'profile_pic' in employee:
                return employee['profile_pic']
    except Exception:
        pass

    try:
        with shelve.open('Admin') as adb:
            admin = adb.get('accounts', {}).get(str(user_id))
            if admin and 'profile_pic' in admin:
                return admin['profile_pic']
    except Exception:
        pass

    return '/static/images/default_profile.png'

@app.route('/api/chats/get-or-create', methods=['GET'])
@login_required
def get_or_create_chat():
    current_user = session.get('employee_id') or session.get('username')  # int or str
    other_user = request.args.get('user_id')

    if not other_user:
        return jsonify({'success': False, 'error': 'Missing user_id'}), 400

    # Normalize the key
    try:
        ids = sorted([str(current_user), str(other_user)], key=int)  # both as strings
    except ValueError:
        return jsonify({'success': False, 'error': 'Invalid user IDs'}), 400

    chat_key = "_".join(ids)

    with shelve.open('ChatDB', writeback=True) as db:
        if chat_key not in db:
            print("Creating new chat:", chat_key)
            db[chat_key] = encrypt_data({
                'participants': ids,
                'messages': {},  # {sender_id: [(time, message)]}
                'last_read': {ids[0]: 0, ids[1]: 0},
                'type': 'individual',
                'calls': []
            })

        return jsonify({'success': True, 'chat_key': chat_key})

# Revised create_chat endpoint
@app.route('/api/create-chat', methods=['POST'])
@login_required
def create_chat_route():
    data = request.get_json()
    user_id = session.get('employee_id') or session.get('username')
    if isinstance(user_id, str) and user_id.isdigit():
        user_id = int(user_id)

    participants = [int(p) for p in data.get('participants', [])]
    chat_type = data.get('type', 'individual')

    if not participants:
        return jsonify({'error': 'No participants specified'}), 400

    if user_id not in participants:
        participants.append(user_id)

    participants = sorted(set(participants))

    chat_key = get_chat_key(participants)

    with shelve.open('ChatDB') as db:
        if chat_key in db:
            return jsonify({'success': True, 'chat_key': chat_key})

        chat_data = {
            'participants': participants,
            'type': chat_type,
            'last_activity': datetime.now().isoformat(),
            'messages': {str(p): [] for p in participants},  # important: each participant has a list
            'read_positions': {str(p): -1 for p in participants}
        }

        if chat_type == 'group':
            chat_data['group_name'] = data.get('group_name', f"Group {datetime.now().strftime('%m%d')}")
            chat_data['group_pic'] = data.get('group_pic', '/static/images/default_group.png')

        db[chat_key] = encrypt_data(chat_data)

    return jsonify({'success': True, 'chat_key': chat_key})

# Revised get_messages endpoint
# Updated get_messages endpoint with pagination
@app.route('/api/chats/messages', methods=['GET'])
@login_required
def get_messages():
    chat_key = request.args.get('key')
    if not chat_key:
        return jsonify({'messages': []})

    try:
        with shelve.open('ChatDB') as db:
            if chat_key not in db:
                return jsonify({'messages': []})

            chat_data = decrypt_data(db[chat_key])
            messages = []

            # Extract messages from chat_data
            if 'messages' in chat_data:
                for sender_id, msg_list in chat_data['messages'].items():
                    username = get_display_name(sender_id)
                    for msg in msg_list:
                        if isinstance(msg, (list, tuple)) and len(msg) >= 2:
                            messages.append({
                                'sender': sender_id,
                                'time': msg[0],
                                'content': msg[1],
                                'username': username
                            })

            return jsonify({
                'messages': sorted(messages, key=lambda x: x['time']),
                'success': True
            })
    except Exception as e:
        print(f"Error in get_messages: {str(e)}")
        return jsonify({'messages': [], 'success': False})

def get_display_name(user_id):
    """
    Return a readable display name for user_id (string or int).
    Tries common shelves/keys used in your project.
    """
    uid = str(user_id)

    # Try User shelves (employees / employee)
    try:
        with shelve.open('User') as udb:
            for key in ('employees', 'employee'):
                emps = udb.get(key, {}) or {}
                # handle keys that might be strings or ints
                for k, v in emps.items():
                    if str(k) == uid:
                        if isinstance(v, dict):
                            return v.get('name') or v.get('username') or v.get('full_name') or uid
                        # object-like
                        name = getattr(v, 'name', None) or getattr(v, 'get_name', lambda: None)()
                        if name:
                            return name
                        return uid
    except Exception:
        pass

    # Try Admin shelves (accounts / admin)
    try:
        with shelve.open('Admin') as adb:
            for key in ('accounts', 'admin'):
                admins = adb.get(key, {}) or {}
                for k, v in admins.items():
                    if str(k) == uid:
                        if isinstance(v, dict):
                            return v.get('username') or v.get('name') or uid
                        name = getattr(v, 'username', None) or getattr(v, 'name', None)
                        if name:
                            return name
                        return uid
    except Exception:
        pass

    # fallback to the id string
    return uid



user_sessions = {}  # sid -> user_id mapping

@socketio.on('connect')
def handle_connect(auth):
    user_id = auth.get('user_id')
    username = auth.get('username')

    print(f'Auth payload: {auth}')  # ✅ Add this for debugging
    if not user_id:
        print("Rejected: No user_id provided")
        return False
    user_sessions[request.sid] = user_id
    print(f"Socket connected: {request.sid} for user {user_id}")


def normalize_messages(chat_data):
    if 'messages' not in chat_data:
        return chat_data

    normalized = {}
    for user_id, messages in chat_data['messages'].items():
        normalized[user_id] = []
        for msg in messages:
            # Handle both formats
            if isinstance(msg, dict) and 'time' in msg and 'content' in msg:
                normalized[user_id].append((msg['time'], msg['content']))
            elif isinstance(msg, (list, tuple)) and len(msg) == 2:
                normalized[user_id].append((msg[0], msg[1]))
            elif isinstance(msg, dict):  # Handle single-item dict case
                for time, content in msg.items():
                    normalized[user_id].append((time, content))

    chat_data['messages'] = normalized
    return chat_data


# Revised WebSocket handler
@socketio.on('send_message')
def handle_send_message(data):
    try:
        chat_key = data['chat_key']
        user_id = user_sessions.get(request.sid)

        if not user_id:
            return {'success': False, 'error': 'Unauthorized'}

        # Use ChatDB consistently
        with shelve.open('ChatDB', writeback=True) as db:
            if chat_key not in db:
                return {'success': False, 'error': 'Chat not found'}

            chat_data = decrypt_data(db[chat_key])
            timestamp = datetime.now().isoformat()

            # Initialize messages structure if needed
            if 'messages' not in chat_data:
                chat_data['messages'] = {}
            if str(user_id) not in chat_data['messages']:
                chat_data['messages'][str(user_id)] = []

            # Store in array format [timestamp, content]
            chat_data['messages'][str(user_id)].append([
                timestamp,
                data['message']
            ])

            db[chat_key] = encrypt_data(chat_data)

            # Emit the new message
            display_name = get_display_name(user_id)
            emit('new_message', {
                'chat_key': chat_key,
                'sender': user_id,
                'time': timestamp,
                'message': data['message'],
                'username': display_name
            }, room=chat_key)

        return {'success': True}
    except Exception as e:
        print(f"Message save error: {str(e)}")
        return {'success': False, 'error': str(e)}

@socketio.on('join_chat')
def handle_join_chat(chat_key):
    join_room(chat_key)


# Search users for chat creation
@app.route('/api/search-users')
def search_users():
    query = request.args.get('q', '').strip().lower()
    results = []

    if not query:
        return jsonify(results)

    with shelve.open('User', 'r') as db:
        employees = db.get('employee', {})
        admins = db.get('admin', {})

        # Search employees
        for eid, emp in employees.items():
            name_match = query in emp.get_name().lower()
            email_match = query in emp.get_email().lower()
            id_match = str(emp.get_ID()) == query

            if name_match or email_match or id_match:
                results.append({
                    'id': emp.get_ID(),
                    'name': emp.get_name(),
                    'username': emp.get_email(),
                    'role': 'employee',
                    'profile_pic': '/static/images/default_profile.png'  # or use actual if stored
                })

        # Search admins
        for uname, adm in admins.items():
            username_match = query in uname.lower()
            name_match = query in adm.get('name', '').lower()  # assuming admin has a 'name' field
            email_match = query in adm.get('email', '').lower()

            if username_match or name_match or email_match:
                results.append({
                    'id': uname,
                    'name': adm.get('name', uname),
                    'username': uname,
                    'role': 'admin',
                    'profile_pic': '/static/images/admin-icon.png'
                })

    return jsonify(results)


# Create new chat
@app.route('/api/create-chat', methods=['POST'])
@login_required
def create_chat():
    data = request.get_json(force=True)

    # normalize current user id as string
    if session.get('is_employee'):
        user_id = str(session['employee_id'])
    elif session.get('is_admin'):
        user_id = str(session['username'])
    else:
        return jsonify({'error': 'Unauthorized'}), 401

    # participants as strings
    participants = [str(p) for p in data.get('participants', []) if str(p)]
    if user_id not in participants:
        participants.append(user_id)
    participants = sorted(set(participants))

    chat_type = data.get('type', 'individual')
    now = datetime.now().isoformat()

    chat_key = get_chat_key(participants)

    with shelve.open('ChatDB') as db:
        if chat_key in db:
            return jsonify({'success': True, 'chat_key': chat_key})

        chat_data = {
            'participants': participants,
            'type': chat_type,                 # 'individual' | 'group' | 'ai'
            'last_activity': now,
            'messages': {pid: [] for pid in participants},    # {sender_id: [(time, content), ...]}
            'read_positions': {pid: -1 for pid in participants}
        }
        if chat_type == 'group':
            chat_data['group_name'] = data.get('group_name', f"Group {datetime.now().strftime('%m%d')}")
            chat_data['group_pic'] = data.get('group_pic', '/static/images/default_group.png')

        db[chat_key] = encrypt_data(chat_data)

    return jsonify({'success': True, 'chat_key': chat_key})


# Initialize AI chat if it doesn't exist
def init_ai_chat():
    with shelve.open('ChatDB') as db:
        if 'ai_chat' not in db:
            ai_chat_data = [
                datetime.now().isoformat(),  # Last activity
                {'ai': []},  # AI messages storage
                [{'user': -1}],  # Read position
                'ai',  # Chat type
                []  # Call history
            ]
            db['ai_chat'] = encrypt_data(ai_chat_data)


@app.route('/api/users/batch', methods=['POST'])
def batch_users():
    user_ids = request.json.get('user_ids', [])
    results = []

    # Search employees
    with shelve.open('User') as db:
        employees = db.get('employees', {})
        for user_id in user_ids:
            emp = employees.get(str(user_id), {})
            if emp:
                results.append({
                    'id': user_id,
                    'name': emp.get('name', f'User {user_id}'),
                    'profile_pic': emp.get('profile_pic', url_for('static', filename='images/default_profile.png')),
                    'type': 'employee'
                })

            # Search admins
            with shelve.open('Admin') as db:
                admins = db.get('accounts', {})
            for user_id in user_ids:
                admin = admins.get(str(user_id), {})
            if admin:
                results.append({
                    'id': user_id,
                    'name': admin.get('username', f'Admin {user_id}'),
                    'profile_pic': url_for('static', filename='images/admin_icon.png'),
                    'type': 'admin'
                })

    return jsonify(results)

@app.route('/api/login', methods=['POST'])
def login():
    try:
        # Check content type
        if request.content_type != 'application/x-www-form-urlencoded':
            return jsonify({
                'success': False,
                'message': 'Invalid content type'
            }), 400

        username = request.form.get('username')
        password = request.form.get('password')

        if not username or not password:
            return jsonify({
                'success': False,
                'message': 'Username and password required'
            }), 400

        with shelve.open('Admin') as db:
            # Check admin accounts first (without ID conversion)
            if username in db.get('accounts', {}):
                stored_hash = db['accounts'][username].encode('utf-8')
                if bcrypt.checkpw(password.encode('utf-8'), stored_hash):
                    session.clear()
                    session['is_admin'] = True
                    session['is_employee'] = False
                    session['username'] = username
                    return jsonify({
                        'success': True,
                        'message': 'Admin login successful',
                        'role': 'admin'
                    })

            # Only attempt employee ID conversion if admin login failed
            try:
                employee_id = int(username)

                # Check employee accounts
                employees = db.get('employees', {})
                if employee_id in employees:
                    creds = employees[employee_id]
                    stored_hash = creds['password'].encode('utf-8')
                    if bcrypt.checkpw(password.encode('utf-8'), stored_hash):
                        session.clear()
                        session['is_employee'] = True
                        session['is_admin'] = False
                        session['employee_id'] = employee_id
                        session['username'] = creds.get('username', str(employee_id))

                        return jsonify({
                            'success': True,
                            'message': 'Employee login successful',
                            'role': 'employee',
                            'employee_id': employee_id,
                            'username': creds.get('username', str(employee_id))
                        })
            except ValueError:
                pass  # Not a numeric employee ID

        return jsonify({
            'success': False,
            'message': 'Invalid credentials'
        }), 401

    except Exception as e:
        print(f"Login error: {str(e)}")
        return jsonify({
            'success': False,
            'message': 'Server error'
        }), 500


@app.route('/reset_password', methods=['POST'])
def reset_password():
    data = request.get_json()
    if not data:
        return jsonify({'success': False, 'error': 'No data provided'}), 400

    username = data.get('username')
    if not username:
        return jsonify({'success': False, 'error': 'Username is required'}), 400

    try:
        employee_id = int(username)
    except ValueError:
        return jsonify({'success': False, 'error': 'Employee ID must be a number'}), 400

    try:
        # Fetch email from User DB
        with shelve.open('User') as user_db:
            employees_user = user_db.get('employee', {})
            if employee_id not in employees_user:
                return jsonify({'success': False, 'error': 'Employee not found in User DB'}), 404

            employee_obj = employees_user[employee_id]
            email = employee_obj.get_email()  # Use class getter

        # Generate new password
        new_password = generate_password()
        hashed_password = bcrypt.hashpw(new_password.encode(), bcrypt.gensalt()).decode()

        # Update Admin DB
        with shelve.open('Admin', writeback=True) as admin_db:
            employees_admin = admin_db.get('employees', {})
            employees_admin[employee_id] = {"password": hashed_password}
            admin_db['employees'] = employees_admin

        # Send email
        if not send_credentials(email, str(employee_id), new_password, is_reset=True):
            app.logger.warning(f"Password reset succeeded but email failed to send for {employee_id}")

        return jsonify({
            'success': True,
            'message': 'Password reset successful. Check your email for the new password.'
        })

    except Exception as e:
        app.logger.error(f"Password reset internal error for {employee_id}: {e}", exc_info=True)
        return jsonify({'success': False, 'error': 'An internal error occurred'}), 500


@app.route('/api/logout', methods=['POST'])
def logout():
    session.clear()
    return jsonify({'success': True, 'message': 'Logged out'})

@app.route('/api/check_session')
def check_session():
    if session.get('is_employee'):
        return jsonify({
            'is_employee': True,
            'is_admin': False,
            'user_id': session.get('employee_id'),
            'username': session.get('username', '')
        })
    elif session.get('is_admin'):
        return jsonify({
            'is_employee': False,
            'is_admin': True,
            'user_id': session.get('username'),  # Admin username as ID
            'username': session.get('username', '')
        })
    else:
        return jsonify({'is_employee': False, 'is_admin': False})

@app.context_processor
def inject_auth_status():
    return {
        'is_admin': session.get('is_admin', False),
        'is_employee': session.get('is_employee', False),
        'employee_id': session.get('employee_id'),
        'username': session.get('username')
    }


@app.route('/api/employee/profile', methods=['GET', 'POST'])
def employee_profile():
    if 'employee_id' not in session:
        return jsonify({'success': False, 'message': 'Not authenticated'}), 401

    emp_id = session['employee_id']

    if request.method == 'GET':
        with shelve.open('Admin') as db:
            employee_data = db.get('employees', {}).get(emp_id, {})
            return jsonify({
                'success': True,
                'username': employee_data.get('username', str(emp_id)),
                'profile_picture': employee_data.get('profile_picture',
                                                     url_for('static', filename='images/default_profile.png'))
            })

    elif request.method == 'POST':
        username = request.form.get('username')
        new_password = request.form.get('password')
        profile_picture = request.files.get('profile_picture')

        with shelve.open('Admin', writeback=True) as db:
            if 'employees' not in db:
                return jsonify({'success': False, 'message': 'Database error'}), 500

            employee_data = db['employees'].get(emp_id, {})

            # Update username
            if username:
                employee_data['username'] = username

            # Update password
            if new_password:
                # Hash the new password
                hashed_password = bcrypt.hashpw(new_password.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')
                employee_data['password'] = hashed_password

            # Update profile picture
            if profile_picture:
                filename = f"profile_{emp_id}.{profile_picture.filename.split('.')[-1]}"
                save_path = os.path.join(current_app.static_folder, 'images', 'profiles', filename)
                profile_picture.save(save_path)
                employee_data['profile_picture'] = url_for('static', filename=f'images/profiles/{filename}')

            # Save back to database
            db['employees'][emp_id] = employee_data
            db.sync()

            return jsonify({
                'success': True,
                'message': 'Profile updated successfully',
                'profile_picture': employee_data.get('profile_picture',
                                                     url_for('static', filename='images/default_profile.png'))
            })

if __name__ == '__main__':
    socketio.run(app, debug=True, allow_unsafe_werkzeug=True)